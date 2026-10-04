import os
import tempfile
import unittest
from pathlib import Path

from rag.extractor import extract_resume_text
from rag.retriever import ResumeRetriever
from rag.vector_store import ResumeVectorStore


def _fake_embed(texts):
    vectors = []
    for text in texts:
        vec = [0.0] * 16
        for char in text.lower():
            vec[ord(char) % 16] += 1.0
        norm = sum(value * value for value in vec) ** 0.5 or 1.0
        vectors.append([value / norm for value in vec])
    return vectors


class ResumeExtractionTests(unittest.TestCase):
    def test_extracts_text_from_docx(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sample.docx"
            doc = Document()
            doc.add_heading("Software Engineer", level=1)
            doc.add_paragraph("Worked on Python APIs and React dashboards.")
            doc.add_paragraph("Built scalable services using Docker and Kubernetes.")
            doc.save(path)

            text = extract_resume_text(path.read_bytes(), path.name)

            self.assertIn("Software Engineer", text)
            self.assertIn("Python APIs", text)
            self.assertIn("Docker", text)


class ResumeIsolationTests(unittest.TestCase):
    def test_sessions_cannot_retrieve_each_others_resume(self):
        from unittest.mock import patch

        with tempfile.TemporaryDirectory() as tmpdir:
            store = ResumeVectorStore(persist_directory=tmpdir)
            retriever = ResumeRetriever(vector_store=store)
            with patch("rag.retriever.embed_texts", side_effect=_fake_embed), patch(
                "rag.vector_store.embed_texts", side_effect=_fake_embed
            ):
                retriever.index_resume("session-a", "Alice built a Kubernetes billing platform in Go.")
                retriever.index_resume("session-b", "Bob taught kindergarten literacy workshops.")

                alice_hits = retriever.retrieve_context("session-a", role="Software Engineer", top_k=3)
                bob_hits = retriever.retrieve_context("session-b", role="Teacher", top_k=3)

            self.assertTrue(alice_hits)
            self.assertTrue(bob_hits)
            self.assertTrue(any("Kubernetes" in hit or "billing" in hit for hit in alice_hits))
            self.assertFalse(any("kindergarten" in hit.lower() for hit in alice_hits))
            self.assertFalse(any("kubernetes" in hit.lower() for hit in bob_hits))
            store.close()


if __name__ == "__main__":
    unittest.main()
