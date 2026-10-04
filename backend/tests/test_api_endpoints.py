from io import BytesIO

from docx import Document


def _docx_bytes(text_heading: str, paragraphs: list[str]) -> bytes:
    doc = Document()
    doc.add_heading(text_heading, level=1)
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_health_and_new_session(client):
    health = client.get("/health")
    assert health.status_code == 200
    assert health.get_json()["status"] == "ok"

    created = client.post("/new_session")
    assert created.status_code == 200
    assert created.get_json()["session_id"]


def test_chat_requires_message(client):
    response = client.post("/chat", json={"session_id": "s1"})
    assert response.status_code == 400


def test_chat_starts_interview_for_known_role(client):
    response = client.post(
        "/chat",
        json={"session_id": "api-role", "message": "I want to practice for a software engineer role"},
    )
    payload = response.get_json()
    assert response.status_code == 200
    assert payload["done"] is False
    assert payload["state_info"]["role_confirmed"] is True
    assert payload["state_info"]["role"] == "Software Engineer"
    assert payload["response"]


def test_chat_routes_vague_answer_to_follow_up(client):
    client.post("/chat", json={"session_id": "api-vague", "message": "software engineer please"})
    response = client.post("/chat", json={"session_id": "api-vague", "message": "I did stuff"})
    payload = response.get_json()
    assert payload["state_info"]["classification"] == "VAGUE"
    assert payload["state_info"]["difficulty"] == "easy"


def test_chat_routes_off_topic_and_out_of_scope(client):
    client.post("/chat", json={"session_id": "api-off", "message": "software engineer"})
    off_topic = client.post("/chat", json={"session_id": "api-off", "message": "Let's talk about pepperoni pizza"})
    assert off_topic.get_json()["state_info"]["classification"] == "OFF_TOPIC"

    client.post("/chat", json={"session_id": "api-scope", "message": "software engineer"})
    out_of_scope = client.post(
        "/chat",
        json={"session_id": "api-scope", "message": "Please write my resume for this job"},
    )
    assert out_of_scope.get_json()["state_info"]["classification"] == "OUT_OF_SCOPE"


def test_resume_upload_endpoint(client):
    payload = {
        "file": (BytesIO(_docx_bytes("Jane Doe", ["Built Python APIs."])), "resume.docx"),
        "session_id": "api-resume",
    }
    response = client.post("/resume/upload", data=payload, content_type="multipart/form-data")
    body = response.get_json()
    assert response.status_code == 200
    assert body["chunks"] >= 1
    assert body["filename"] == "resume.docx"


def test_uploaded_resume_reaches_question_prompt(client, monkeypatch):
    prompts = []

    def capture_chat(system_prompt, history, user_message, **kwargs):
        prompts.append(user_message)
        return "Tell me about the payment platform you built."

    monkeypatch.setattr("graph.nodes.chat_completion", capture_chat)
    payload = {
        "file": (
            BytesIO(_docx_bytes("Resume", ["Built a Kubernetes payment platform in Go."])),
            "resume.docx",
        ),
        "session_id": "api-rag",
    }
    upload = client.post("/resume/upload", data=payload, content_type="multipart/form-data")
    assert upload.status_code == 200

    response = client.post(
        "/chat",
        json={"session_id": "api-rag", "message": "I want to practice for software engineer"},
    )
    assert response.status_code == 200
    assert any("Kubernetes" in prompt or "payment" in prompt for prompt in prompts)


def test_feedback_endpoint_with_transcript(client, llm_mocks):
    response = client.post(
        "/feedback",
        json={
            "session_id": "feedback-only",
            "role": "Software Engineer",
            "transcript": [
                {"role": "assistant", "content": "Tell me about a project."},
                {"role": "user", "content": "I built an API gateway."},
            ],
        },
    )
    assert response.status_code == 200
    feedback = response.get_json()["feedback"]
    assert "overallImpression" in feedback
