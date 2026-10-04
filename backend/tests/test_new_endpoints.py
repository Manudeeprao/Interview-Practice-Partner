"""Tests for the new endpoints, validation, error shapes, and rate limiting."""

import hashlib
import io
import math

import pytest
from docx import Document

from app import RateLimiter, create_app


def _fake_embed(texts, batch_size=32):
    vectors = []
    for text in texts:
        digest = hashlib.md5(text.encode("utf-8")).digest()
        vec = [b / 255.0 for b in digest[:16]]
        norm = math.sqrt(sum(x * x for x in vec)) or 1.0
        vectors.append([x / norm for x in vec])
    return vectors


@pytest.fixture
def no_model(monkeypatch):
    monkeypatch.setattr("rag.vector_store.embed_texts", _fake_embed)


def _pdf_bytes(text: str) -> bytes:
    # Minimal valid PDF containing text; pypdf can extract it.
    body = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    parts = [
        b"%PDF-1.4\n",
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
        b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n",
        b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >> endobj\n",
        b"4 0 obj << /Length " + str(len(body)).encode() + b" >> stream\n" + body + b"\nendstream endobj\n",
        b"5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n",
    ]
    offset = 0
    xref = []
    out = b""
    for i, part in enumerate(parts):
        if i > 0:  # skip the %PDF header; object offsets are 1-based
            xref.append(offset)
        out += part
        offset += len(part)
    xref_pos = offset
    out += b"xref\n0 6\n0000000000 65535 f \n"
    for pos in xref:
        out += f"{pos:010d} 00000 n \n".encode()
    out += b"trailer << /Size 6 /Root 1 0 R >>\nstartxref\n" + str(xref_pos).encode() + b"\n%%EOF"
    return out


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Built the Zebracorn platform.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Error shapes + validation
# ---------------------------------------------------------------------------
def test_chat_error_shape_has_code(client):
    response = client.post("/chat", json={"session_id": "s1"})
    assert response.status_code == 400
    payload = response.get_json()
    assert payload["error"]
    assert payload["code"] == "missing_message"


def test_chat_rejects_invalid_session_id(client):
    response = client.post("/chat", json={"session_id": "../../etc", "message": "hi"})
    assert response.status_code == 400
    assert response.get_json()["code"] == "invalid_session_id"


def test_chat_rejects_oversized_message(client):
    response = client.post(
        "/chat", json={"session_id": "s1", "message": "x" * 5000}
    )
    assert response.status_code == 400
    assert response.get_json()["code"] == "message_too_long"


def test_get_session_roundtrip(client):
    created = client.post("/new_session").get_json()
    sid = created["session_id"]
    client.post("/chat", json={"session_id": sid, "message": "software engineer"})

    response = client.get(f"/session/{sid}")
    assert response.status_code == 200
    payload = response.get_json()
    assert payload["session_id"] == sid
    assert any(m["role"] == "user" for m in payload["history"])
    assert payload["state_info"]["role_confirmed"] is True


def test_get_session_404(client):
    response = client.get("/session/does-not-exist")
    assert response.status_code == 404
    assert response.get_json()["code"] == "session_not_found"


def test_new_session_resets_existing(client, no_model):
    created = client.post("/new_session").get_json()
    sid = created["session_id"]

    data = {"file": (io.BytesIO(_pdf_bytes("Zebracorn project")), "resume.pdf")}
    upload = client.post(
        "/resume/upload",
        data={**data, "session_id": sid},
        content_type="multipart/form-data",
    )
    assert upload.status_code == 200
    status = client.get(f"/resume/status?session_id={sid}").get_json()
    assert status["uploaded"] is True

    reset = client.post("/new_session", json={"session_id": sid})
    assert reset.status_code == 200
    assert reset.get_json()["session_id"] == sid

    status = client.get(f"/resume/status?session_id={sid}").get_json()
    assert status["uploaded"] is False
    assert status["chunks"] == 0


# ---------------------------------------------------------------------------
# Resume endpoints
# ---------------------------------------------------------------------------
def test_resume_status_empty(client):
    created = client.post("/new_session").get_json()
    status = client.get(f"/resume/status?session_id={created['session_id']}").get_json()
    assert status["uploaded"] is False
    assert status["chunks"] == 0


def test_resume_upload_and_delete(client, no_model):
    created = client.post("/new_session").get_json()
    sid = created["session_id"]

    upload = client.post(
        "/resume/upload",
        data={
            "file": (io.BytesIO(_docx_bytes()), "resume.docx"),
            "session_id": sid,
        },
        content_type="multipart/form-data",
    )
    assert upload.status_code == 200
    payload = upload.get_json()
    assert payload["chunks"] > 0

    status = client.get(f"/resume/status?session_id={sid}").get_json()
    assert status["uploaded"] is True
    assert status["filename"] == "resume.docx"

    deleted = client.delete(f"/resume/{sid}")
    assert deleted.status_code == 200
    assert deleted.get_json()["deleted"] is True

    status = client.get(f"/resume/status?session_id={sid}").get_json()
    assert status["uploaded"] is False
    assert status["chunks"] == 0


def test_resume_upload_rejects_bad_content(client):
    created = client.post("/new_session").get_json()
    sid = created["session_id"]
    response = client.post(
        "/resume/upload",
        data={
            "file": (io.BytesIO(b"definitely not a pdf"), "resume.pdf"),
            "session_id": sid,
        },
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert response.get_json()["code"] == "invalid_file_content"


def test_resume_upload_rejects_bad_extension(client):
    created = client.post("/new_session").get_json()
    sid = created["session_id"]
    response = client.post(
        "/resume/upload",
        data={
            "file": (io.BytesIO(b"hello"), "resume.txt"),
            "session_id": sid,
        },
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert response.get_json()["code"] == "unsupported_file_type"


# ---------------------------------------------------------------------------
# Rate limiting
# ---------------------------------------------------------------------------
def test_rate_limiter_unit():
    limiter = RateLimiter(max_requests=2, window_seconds=60)
    assert limiter.allow("k")[0] is True
    assert limiter.allow("k")[0] is True
    allowed, retry_after = limiter.allow("k")
    assert allowed is False
    assert retry_after > 0


def test_chat_rate_limit_429(tmp_path, monkeypatch, memory_graph, session_store):
    monkeypatch.setenv("RATE_LIMIT_CHAT_PER_MIN", "1")
    monkeypatch.setenv("RATE_LIMIT_UPLOAD_PER_MIN", "100")
    monkeypatch.setattr("rag.vector_store.embed_texts", _fake_embed)
    from rag.retriever import ResumeRetriever
    from rag.vector_store import ResumeVectorStore

    retriever = ResumeRetriever(
        vector_store=ResumeVectorStore(persist_directory=str(tmp_path / "chroma2"))
    )
    flask_app = create_app(
        session_store=session_store,
        interview_graph=memory_graph,
        resume_retriever=retriever,
    )
    test_client = flask_app.test_client()

    first = test_client.post("/chat", json={"session_id": "rl", "message": "software engineer"})
    assert first.status_code == 200
    second = test_client.post("/chat", json={"session_id": "rl", "message": "hello again"})
    assert second.status_code == 429
    assert second.get_json()["code"] == "rate_limited"
